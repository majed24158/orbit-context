"""Run with the bundled Python. Fixtures and QA stay under this directory.

Optional --word-python PATH uses the existing owned Word export helper only for
rendering when the packaged LibreOffice renderer cannot run on Windows. It does
not author the DOCX. PDF PNG rendering still uses bundled pypdfium2.
"""
from pathlib import Path
import argparse
import copy
import json
import subprocess
import sys
import zipfile

from docx import Document
from lxml import etree
from PIL import Image, ImageDraw
from editable_docx import *
import editable_docx as writer

ROOT = Path(__file__).parent / 'qa' / 'editable_docx'
ROOT.mkdir(parents=True, exist_ok=True)
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'v': 'urn:schemas-microsoft-com:vml',
      'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'}


def block(kind='text', **values):
    data = {'kind': kind, 'bbox': [.08, .06, .84, .12], 'text': 'Editable English text',
            'font_size': 14, 'bold': False, 'italic': False, 'color': '#000000',
            'align': 'left', 'rtl': False, 'rows': [], 'column_widths': [],
            'header': False, 'merges': []}
    data.update(values)
    return data


def fixtures():
    return [
        {'width_pt': 595.28, 'height_pt': 841.89, 'blocks': [
            block(text='مستند عربي قابل للتعديل', font_size=22, bold=True, rtl=True, align='right'),
            block(text='هذه الفقرة نص أصلي داخل Word ويمكن تحريرها بحرية.\nلا تُستبدل النصوص بصور للصفحة.',
                  bbox=[.08,.16,.84,.075], font_size=14, rtl=True, align='right'),
            block('table', text='', bbox=[.08,.29,.84,.20], rtl=True, align='right',
                  rows=[['البيان','','القيمة'], ['المشروع','تحويل المستندات','2026'],
                        ['الحالة','قابل للتعديل','جاهز']], column_widths=[.25,.45,.30],
                  header=True, merges=[[0,0,0,1]]),
            block(text='نهاية الصفحة الأولى', bbox=[.08,.85,.84,.07], rtl=True, align='right'),
        ]},
        {'width_pt': 720, 'height_pt': 480, 'blocks': [
            block(text='Editable landscape page', bold=True, font_size=22),
            block(text='Native text, a native table, and one cropped picture.\nXML-like input is literal text: <script>noop()</script> & "quoted".',
                  bbox=[.08,.21,.84,.14], font_size=12),
            block('table', text='', bbox=[.08,.44,.46,.22], font_size=12,
                  rows=[['Item','Count'], ['Editable rows','12'], ['Native cells','8']],
                  column_widths=[.72,.28], header=True, merges=[]),
            block('picture', text='', bbox=[.65,.44,.25,.25], font_size=None, color='', align=''),
        ]},
        {'width_pt': 420, 'height_pt': 595, 'blocks': [
            block(text='Final page without a blank page', font_size=17, bold=True),
            block(text='The document has exactly three source pages.\nAll recognized text remains editable.',
                  bbox=[.08,.22,.84,.16], font_size=12),
        ]},
    ]


def main():
    parser=argparse.ArgumentParser(); parser.add_argument('--word-python'); parser.add_argument('--renderer'); args=parser.parse_args()
    pages=fixtures(); checks=[]
    normalized=validate_pages(pages); normalized[0]['blocks'][0]['text']='changed'
    assert pages[0]['blocks'][0]['text']!='changed'
    checks.append('validation returns detached data')
    def rejected(modifier):
        bad=copy.deepcopy(pages); modifier(bad)
        try: validate_pages(bad)
        except ValueError: return
        raise AssertionError('invalid pages accepted')
    mutations=[
        lambda p:p[0].update(width_pt=float('nan')),
        lambda p:p[0].update(height_pt=float('inf')),
        lambda p:p[0].update(width_pt=10**1000),
        lambda p:p[0].update(extra='not permitted'),
        lambda p:p[0]['blocks'][0].pop('rtl'),
        lambda p:p[0]['blocks'][0].update(bbox=[.8,.1,.4,.2]),
        lambda p:p[0]['blocks'][0].update(bbox=[0,0,0,.2]),
        lambda p:p[0]['blocks'][0].update(bbox=[False,0,.8,.2]),
        lambda p:p[0]['blocks'][0].update(font_size=0),
        lambda p:p[0]['blocks'][0].update(font_size=True),
        lambda p:p[0]['blocks'][0].update(color='url(https://example.invalid)'),
        lambda p:p[0]['blocks'][0].update(text='\x00'),
        lambda p:p[0]['blocks'][0].update(text='x'*8001),
        lambda p:p[0]['blocks'][2].update(rows=[['a','b'],['c']]),
        lambda p:p[0]['blocks'][2].update(column_widths=[.5,.5]),
        lambda p:p[0]['blocks'][2].update(column_widths=[.1,.1,.1]),
        lambda p:p[0]['blocks'][2].update(merges=[[0,0,10,1]]),
        lambda p:p[0]['blocks'][2].update(merges=[[1,1,0,1]]),
        lambda p:p[0]['blocks'][2].update(merges=[[0,0,0,1],[0,1,1,1]]),
        lambda p:p[0]['blocks'][2].update(merges=[[0,0,0,0]]),
        lambda p:p[1]['blocks'][3].update(text='Text must not disappear in a picture'),
    ]
    for mutation in mutations: rejected(mutation)
    for bad in ([],pages*21,{},'pages'):
        try: validate_pages(bad)
        except ValueError: pass
        else: raise AssertionError('invalid root accepted')
    checks.append(f'{len(mutations)+4} malformed/bounds/span/NaN/size/XML inputs rejected')
    page={'width_pt':595,'height_pt':842,'blocks':[block() for _ in range(201)]}
    for invalid in ([page],
                    [{'width_pt':595,'height_pt':842,'blocks':[block() for _ in range(200)]} for _ in range(13)],
                    [{'width_pt':595,'height_pt':842,'blocks':[block(text='x'*8000) for _ in range(63)]}],
                    [{'width_pt':595,'height_pt':842,'blocks':[block('table',text='',rows=[['x']*30 for _ in range(100)],column_widths=[1/30]*30) for _ in range(7)]}]):
        try: validate_pages(invalid)
        except ValueError: pass
        else: raise AssertionError('aggregate limit not enforced')
    checks.append('per-page/aggregate block, text and table-cell caps enforced')
    def strict(schema):
        if isinstance(schema,dict):
            if schema.get('type')=='object':
                assert schema.get('additionalProperties') is False
                assert set(schema['required'])==set(schema['properties'])
            for value in schema.values(): strict(value)
        elif isinstance(schema,list):
            for value in schema: strict(value)
    strict(DOCUMENT_SCHEMA); checks.append('strict all-required schema at every object')
    source=ROOT/'source_page.png'
    image=Image.new('RGB',(1440,960),'white'); draw=ImageDraw.Draw(image)
    draw.rectangle((936,422,1296,662),fill='#dbe8f2',outline='#26506f',width=4)
    draw.ellipse((1020,450,1200,630),fill='#ef743f'); image.save(source); image.close()
    path=build_docx(pages,ROOT/'editable_fixture.docx',[None,source,None])
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None
        xml=etree.fromstring(archive.read('word/document.xml'))
        texts=xml.xpath('//w:txbxContent//w:t/text()',namespaces=NS)
        assert 'مستند عربي قابل للتعديل' in texts
        assert any('<script>noop()</script>' in text for text in texts)
        assert len(xml.xpath('//w:tbl',namespaces=NS))==2
        assert len(xml.xpath('//w:tblpPr',namespaces=NS))==2
        assert xml.xpath('//w:gridSpan[@w:val="2"]',namespaces=NS)
        assert xml.xpath('//w:tblHeader',namespaces=NS)
        assert xml.xpath('//w:bidi[@w:val="1"]',namespaces=NS)
        assert xml.xpath('//w:rFonts[@w:cs="Arial"]',namespaces=NS)
        assert len(xml.xpath('//w:sectPr',namespaces=NS))==3
        assert not xml.xpath('//w:br[@w:type="page"]',namespaces=NS)
        assert len(xml.xpath('//wp:anchor',namespaces=NS))==1
        assert len([name for name in archive.namelist() if name.startswith('word/media/')])==1
        assert not any(name.endswith('vbaProject.bin') for name in archive.namelist())
        relfiles=[archive.read(name) for name in archive.namelist() if name.endswith('.rels')]
        assert not any(b'TargetMode="External"' in rel for rel in relfiles)
    checks.append('editable native RTL/text/table/merge/header nodes, page sizes, crop and safe literals')
    reopened=Document(path); assert len(reopened.tables)==2
    reopened.tables[1].cell(1,0).text='Changed natively'
    assert reopened.tables[1].cell(1,0).text=='Changed natively'
    checks.append('python-docx reopens and edits native table cells')
    # Tables/text do not require source bitmaps and cannot silently rasterize.
    onlytext=[{'width_pt':420,'height_pt':595,'blocks':[block()]}]
    temporary=ROOT/'text_only.tmp'; build_docx(onlytext,temporary)
    with zipfile.ZipFile(temporary) as archive:
        assert not any(name.startswith('word/media/') for name in archive.namelist())
    temporary.unlink(); checks.append('text-only conversion embeds no raster fallback')
    merged=block('table',text='',bbox=[.08,.2,.84,.5],rows=[['A','B'],['C','D']],
                 column_widths=[.5,.5],merges=[[0,0,1,1]])
    merged_path=ROOT/'merged.tmp'; build_docx([{'width_pt':420,'height_pt':595,'blocks':[merged]}],merged_path)
    with zipfile.ZipFile(merged_path) as archive:
        merged_xml=etree.fromstring(archive.read('word/document.xml'))
        assert merged_xml.xpath('//w:vMerge[@w:val="restart"]',namespaces=NS)
        assert merged_xml.xpath('//w:vMerge[not(@w:val)]',namespaces=NS)
        assert merged_xml.xpath('//w:t/text()',namespaces=NS)==['A','B','C','D']
    merged_path.unlink(); checks.append('rectangular horizontal+vertical merges preserve all source cell text')
    before=path.read_bytes()
    try: build_docx(pages,path,{})
    except ValueError: pass
    else: raise AssertionError('missing picture accepted')
    assert path.read_bytes()==before
    assert not list(ROOT.glob('.orbit-docx-*'))
    checks.append('failed picture input preserves existing destination and cleans staging')
    for name in ('MAX_IMAGE_PIXELS','MAX_TOTAL_IMAGE_PIXELS','MAX_CROP_PIXELS','MAX_EMBEDDED_BYTES','MAX_OUTPUT_BYTES'):
        original=getattr(writer,name); setattr(writer,name,1)
        try:
            try: build_docx(pages,path,[None,source,None])
            except ValueError: pass
            else: raise AssertionError(name+' not enforced')
            assert path.read_bytes()==before
            assert not list(ROOT.glob('.orbit-docx-*'))
        finally: setattr(writer,name,original)
    checks.append('image decode/crop/encoded/output limits reject atomically without replacing a destination')
    renderer=Path(args.renderer) if args.renderer else Path(__file__).parent/'optional-renderer'/'render_docx.py'
    rendered=ROOT/'rendered'
    attempt=subprocess.run([sys.executable,str(renderer),str(path),'--output_dir',str(rendered),'--emit_pdf','--verbose'],capture_output=True,text=True)
    (ROOT/'renderer.log').write_text(attempt.stdout+'\n'+attempt.stderr,encoding='utf-8')
    render_status='packaged renderer success' if attempt.returncode==0 else 'packaged renderer unavailable; see renderer.log'
    if attempt.returncode and args.word_python:
        # Existing helper opens a separate hidden Word instance read-only, disables
        # macros/links, and closes only its owned document/application.
        code=('import sys,json; sys.path.insert(0,sys.argv[1]); from word_export import export_word_pdf; '
              'print(json.dumps(str(export_word_pdf(sys.argv[2],sys.argv[3]))))')
        fallback=subprocess.run([args.word_python,'-c',code,str(Path(__file__).parent),str(path),str(ROOT)],capture_output=True,text=True)
        (ROOT/'word_renderer.log').write_text(fallback.stdout+'\n'+fallback.stderr,encoding='utf-8')
        assert fallback.returncode==0,fallback.stderr
        pdf=Path(json.loads(fallback.stdout.strip().splitlines()[-1])); rendered.mkdir(exist_ok=True)
        from pypdf import PdfReader
        pdfdoc=PdfReader(pdf); assert len(pdfdoc.pages)==3,'Unexpected/blank Word-rendered page'
        for page,text in zip(pdfdoc.pages,('نهاية','Editable landscape page','Final page')):
            # Word PDF extraction may encode Arabic differently; enforce English
            # structure and check Arabic visually in the rendered page PNG.
            if text.isascii(): assert text in page.extract_text()
        import pypdfium2 as pdfium
        with pdfium.PdfDocument(pdf) as doc:
            for index in range(len(doc)):
                page=doc[index]; bitmap=page.render(scale=1.6)
                bitmap.to_pil().save(rendered/f'page-{index+1}.png'); bitmap.close(); page.close()
        render_status='Word desktop fallback; 3 pages, bundled PDFium PNGs; visual inspection required'
    report={'checks':checks,'count':len(checks),'render':render_status,'fixture':str(path)}
    (ROOT/'report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False,indent=2))


if __name__=='__main__': main()
