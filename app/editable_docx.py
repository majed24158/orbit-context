"""Bounded, deterministic reconstruction of editable Word pages.

This module consumes data, never Python/XML supplied by an OCR/model. Text is
stored in native w:p/w:r nodes inside VML text boxes; tables are native w:tbl
objects positioned with tblpPr. Only explicit picture blocks embed raster crops.
Word desktop is the layout authority; web/third-party viewers may move floats.

Coordinates are page-relative fractions. Table rows/columns use logical order;
rtl=True renders the first logical column at the right. Merges have inclusive
zero-based corners and preserve all nonempty covered-cell text. Reconstruction
is approximate: font matching, OCR, text fit, and reading order remain upstream
responsibilities. Text boxes retain their supplied bounds; oversized content
can clip and needs upstream block/font correction plus Word-rendered review.
"""
from __future__ import annotations

import copy
import io
import math
import os
from pathlib import Path
import re
import tempfile
from collections.abc import Mapping, Sequence

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from lxml import etree
from PIL import Image, ImageOps

MAX_PAGES = 60
MAX_BLOCKS_PER_PAGE = 200
MAX_BLOCKS = 2500
MAX_TEXT_CHARS = 500_000
MAX_BLOCK_CHARS = 8000
MAX_CELL_CHARS = 2000
MAX_TABLE_ROWS = 100
MAX_TABLE_COLUMNS = 30
MAX_TABLE_CELLS = 20_000
MAX_IMAGE_PIXELS = 40_000_000
MAX_TOTAL_IMAGE_PIXELS = 120_000_000
MAX_IMAGE_BYTES = 40_000_000
MAX_CROP_PIXELS = 80_000_000
MAX_EMBEDDED_BYTES = 60_000_000
MAX_OUTPUT_BYTES = 100_000_000

_BLOCK_FIELDS = ('kind', 'bbox', 'text', 'font_size', 'bold', 'italic', 'color',
                 'align', 'rtl', 'rows', 'column_widths', 'header', 'merges')
_ALIGNMENTS = ('left', 'right', 'center', 'justify')
_BAD_XML = re.compile('[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]')
_COLOR = re.compile(r'#[0-9a-fA-F]{6}\Z')
_NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
       'v': 'urn:schemas-microsoft-com:vml',
       'o': 'urn:schemas-microsoft-com:office:office',
       'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'}

BLOCK_SCHEMA = {
    'type': 'object', 'additionalProperties': False,
    'required': list(_BLOCK_FIELDS),
    'properties': {
        'kind': {'type': 'string', 'enum': ['text', 'table', 'picture']},
        'bbox': {'type': 'array', 'minItems': 4, 'maxItems': 4,
                 'items': {'type': 'number', 'minimum': 0, 'maximum': 1}},
        'text': {'type': 'string', 'maxLength': MAX_BLOCK_CHARS},
        'font_size': {'type': ['number', 'null'], 'minimum': 6, 'maximum': 96},
        'bold': {'type': 'boolean'}, 'italic': {'type': 'boolean'},
        'color': {'type': 'string', 'pattern': r'^(#[0-9a-fA-F]{6})?$'},
        'align': {'type': 'string', 'enum': ['', *_ALIGNMENTS]},
        'rtl': {'type': 'boolean'},
        'rows': {'type': 'array', 'maxItems': MAX_TABLE_ROWS, 'items': {
            'type': 'array', 'maxItems': MAX_TABLE_COLUMNS,
            'items': {'type': 'string', 'maxLength': MAX_CELL_CHARS}}},
        'column_widths': {'type': 'array', 'maxItems': MAX_TABLE_COLUMNS,
                          'items': {'type': 'number', 'exclusiveMinimum': 0, 'maximum': 1}},
        'header': {'type': 'boolean'},
        'merges': {'type': 'array', 'maxItems': MAX_TABLE_ROWS * MAX_TABLE_COLUMNS // 2,
                   'items': {'type': 'array', 'minItems': 4, 'maxItems': 4,
                             'items': {'type': 'integer', 'minimum': 0}}},
    },
}
PAGE_SCHEMA = {'type': 'object', 'additionalProperties': False,
               'required': ['width_pt', 'height_pt', 'blocks'], 'properties': {
                   'width_pt': {'type': 'number', 'minimum': 72, 'maximum': 1584},
                   'height_pt': {'type': 'number', 'minimum': 72, 'maximum': 1584},
                   'blocks': {'type': 'array', 'maxItems': MAX_BLOCKS_PER_PAGE,
                              'items': BLOCK_SCHEMA}}}
DOCUMENT_SCHEMA = {'type': 'object', 'additionalProperties': False,
                   'required': ['pages'], 'properties': {
                       'pages': {'type': 'array', 'minItems': 1, 'maxItems': MAX_PAGES,
                                 'items': PAGE_SCHEMA}}}


def _fail(where, message):
    raise ValueError(f'{where}: {message}')


def _number(value, where, low, high):
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        _fail(where, 'expected a finite number')
    if not low <= value <= high or not math.isfinite(value):
        _fail(where, f'expected a finite number from {low} to {high}')
    return float(value)


def _text(value, where, cap):
    if not isinstance(value, str) or len(value) > cap or _BAD_XML.search(value):
        _fail(where, f'expected XML-safe text of at most {cap} characters')
    return value


def validate_pages(pages):
    """Validate all data before filesystem work; return a detached plain copy.

    Extra/missing keys, NaN/Infinity, ragged tables, overlapping/out-of-range
    merges, invalid colors/controls and excessive aggregate sizes are rejected.
    Schema alone cannot enforce bbox sums, table rectangles, or merge overlap.
    """
    if not isinstance(pages, list) or not 1 <= len(pages) <= MAX_PAGES:
        _fail('pages', f'expected 1 to {MAX_PAGES} pages')
    total_blocks = total_chars = total_cells = 0
    for pi, page in enumerate(pages):
        where = f'pages[{pi}]'
        if not isinstance(page, dict) or set(page) != {'width_pt', 'height_pt', 'blocks'}:
            _fail(where, 'page keys must be width_pt, height_pt, blocks')
        width = _number(page['width_pt'], where + '.width_pt', 72, 1584)
        height = _number(page['height_pt'], where + '.height_pt', 72, 1584)
        blocks = page['blocks']
        if not isinstance(blocks, list) or len(blocks) > MAX_BLOCKS_PER_PAGE:
            _fail(where, 'too many blocks or blocks is not a list')
        total_blocks += len(blocks)
        if total_blocks > MAX_BLOCKS:
            _fail(where, 'document has too many blocks')
        for bi, block in enumerate(blocks):
            pos = f'{where}.blocks[{bi}]'
            if not isinstance(block, dict) or set(block) != set(_BLOCK_FIELDS):
                _fail(pos, 'block has missing or unexpected keys')
            if block['kind'] not in ('text', 'table', 'picture'):
                _fail(pos, 'unknown block kind')
            bbox = block['bbox']
            if not isinstance(bbox, list) or len(bbox) != 4:
                _fail(pos + '.bbox', 'expected [left, top, width, height]')
            left, top, bw, bh = [_number(v, pos + '.bbox', 0, 1) for v in bbox]
            if bw * width < 1 or bh * height < 1 or left + bw > 1 + 1e-8 or top + bh > 1 + 1e-8:
                _fail(pos + '.bbox', 'box must have positive physical size and stay inside page')
            for field in ('bold', 'italic', 'rtl', 'header'):
                if type(block[field]) is not bool:
                    _fail(pos + '.' + field, 'expected boolean')
            text = _text(block['text'], pos + '.text', MAX_BLOCK_CHARS)
            total_chars += len(text)
            if block['font_size'] is not None:
                _number(block['font_size'], pos + '.font_size', 6, 96)
            if not isinstance(block['color'], str) or (block['color'] and not _COLOR.fullmatch(block['color'])):
                _fail(pos + '.color', 'expected #RRGGBB or empty picture color')
            if block['align'] not in ('', *_ALIGNMENTS):
                _fail(pos + '.align', 'unknown paragraph alignment')
            if block['kind'] != 'picture':
                if block['font_size'] is None or not block['color'] or not block['align']:
                    _fail(pos, 'text/table needs font_size, color and align')
            for field in ('rows', 'column_widths', 'merges'):
                if not isinstance(block[field], list):
                    _fail(pos + '.' + field, 'expected list')
            if block['kind'] != 'table':
                if block['rows'] or block['column_widths'] or block['merges'] or block['header']:
                    _fail(pos, 'only table blocks may contain rows, widths, merges or a header')
                if block['kind'] == 'text' and not text.strip():
                    _fail(pos, 'text blocks need nonempty text')
                if block['kind'] == 'picture' and text:
                    _fail(pos, 'picture text must be empty; recognized text must be a text block')
            else:
                rows = block['rows']
                if text or not 1 <= len(rows) <= MAX_TABLE_ROWS or not isinstance(rows[0], list):
                    _fail(pos, 'table needs rows and an empty text field')
                columns = len(rows[0])
                if not 1 <= columns <= MAX_TABLE_COLUMNS:
                    _fail(pos, 'table column count is out of range')
                total_cells += len(rows) * columns
                if total_cells > MAX_TABLE_CELLS:
                    _fail(pos, 'document has too many table cells')
                for ri, row in enumerate(rows):
                    if not isinstance(row, list) or len(row) != columns:
                        _fail(pos, 'table rows must form a rectangle')
                    for ci, cell in enumerate(row):
                        total_chars += len(_text(cell, f'{pos}.rows[{ri}][{ci}]', MAX_CELL_CHARS))
                widths = block['column_widths']
                if len(widths) != columns:
                    _fail(pos, 'column_widths must have one positive fraction for each column')
                normalized = [_number(v, pos + '.column_widths', .0001, 1) for v in widths]
                if abs(sum(normalized) - 1) > .01:
                    _fail(pos, 'column width fractions must sum to 1 (tolerance 0.01)')
                covered = set()
                if len(block['merges']) > len(rows) * columns // 2:
                    _fail(pos, 'too many merges')
                for merge in block['merges']:
                    if not isinstance(merge, list) or len(merge) != 4 or any(type(v) is not int for v in merge):
                        _fail(pos, 'merges need four integer inclusive indices')
                    r0, c0, r1, c1 = merge
                    if not (0 <= r0 <= r1 < len(rows) and 0 <= c0 <= c1 < columns) or (r0 == r1 and c0 == c1):
                        _fail(pos, 'merge is outside table, reversed, or a single cell')
                    span = {(r, c) for r in range(r0, r1 + 1) for c in range(c0, c1 + 1)}
                    if covered & span:
                        _fail(pos, 'table merges overlap')
                    covered |= span
            if total_chars > MAX_TEXT_CHARS:
                _fail(pos, 'document has too much text')
    return copy.deepcopy(pages)


def _node(tag, **attrs):
    prefix, local = tag.split(':')
    node = etree.Element('{' + _NS[prefix] + '}' + local)
    for name, value in attrs.items():
        if ':' in name:
            p, n = name.split(':'); name = '{' + _NS[p] + '}' + n
        node.set(name, str(value))
    return node


def _w_property(parent, tag, **attrs):
    element = OxmlElement('w:' + tag)
    for key, value in attrs.items():
        element.set(qn('w:' + key), str(value))
    parent.append(element)
    return element


def _ordered(parent, names):
    """Maintain OOXML property order when adding uncommon native properties."""
    order = {qn('w:' + name): i for i, name in enumerate(names.split())}
    children = list(parent)
    children.sort(key=lambda child: order.get(child.tag, 999))
    parent[:] = children


def _format_paragraph(paragraph, block):
    fmt = paragraph.paragraph_format
    fmt.space_before = fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0
    fmt.keep_together = False
    fmt.keep_with_next = False
    fmt.widow_control = False
    # Word mirrors left/right justification for bidi paragraphs. The data
    # contract uses physical page alignment, so compensate before writing jc.
    alignment = block['align']
    if block['rtl'] and alignment in ('left', 'right'):
        alignment = 'right' if alignment == 'left' else 'left'
    paragraph.alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'right': WD_ALIGN_PARAGRAPH.RIGHT,
                           'center': WD_ALIGN_PARAGRAPH.CENTER, 'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}[alignment]
    _w_property(paragraph._p.get_or_add_pPr(), 'bidi', val='1' if block['rtl'] else '0')
    _w_property(paragraph._p.get_or_add_pPr(), 'snapToGrid', val='0')
    for run in paragraph.runs:
        run.font.name = 'Arial'
        run.font.size = Pt(block['font_size'])
        run.bold = block['bold']; run.italic = block['italic']
        run.font.color.rgb = RGBColor.from_string(block['color'][1:])
        props = run._r.get_or_add_rPr()
        fonts = props.find(qn('w:rFonts'))
        for script in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            fonts.set(qn('w:' + script), 'Arial')
        _w_property(props, 'szCs', val=round(block['font_size'] * 2))
        _w_property(props, 'rtl', val='1' if block['rtl'] else '0')
        if block['bold']: _w_property(props, 'bCs')
        if block['italic']: _w_property(props, 'iCs')
        _w_property(props, 'lang', val='ar-SA' if block['rtl'] else 'en-US', bidi='ar-SA')
        _ordered(props, 'rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs em lang')
    _ordered(paragraph._p.get_or_add_pPr(), 'pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr')


def _anchor_paragraph(document):
    paragraph = document.add_paragraph()
    fmt = paragraph.paragraph_format
    fmt.space_before = fmt.space_after = Pt(0)
    fmt.line_spacing = Pt(1)
    paragraph.add_run().font.size = Pt(1)
    _w_property(paragraph._p.get_or_add_pPr(), 'snapToGrid', val='0')
    return paragraph


def _textbox(anchor, block, box, shape_id):
    from docx.text.paragraph import Paragraph
    left, top, width, height = box
    pict = OxmlElement('w:pict')
    shape_type = _node('v:shapetype', id=f'_orbit_text_type_{shape_id}', coordsize='21600,21600',
                       **{'o:spt': '202', 'path': 'm,l,21600r21600,l21600,xe'})
    shape_type.append(_node('v:stroke', joinstyle='miter'))
    shape_type.append(_node('v:path', gradientshapeok='t', **{'o:connecttype': 'rect'}))
    pict.append(shape_type)
    shape = _node('v:shape', id=f'OrbitText{shape_id}', type=f'#_orbit_text_type_{shape_id}',
                  style=(f'position:absolute;margin-left:{left:.3f}pt;margin-top:{top:.3f}pt;'
                         f'width:{width:.3f}pt;height:{height:.3f}pt;z-index:{shape_id};'
                         'mso-position-horizontal-relative:page;mso-position-vertical-relative:page;'
                         'mso-wrap-style:square'),
                  filled='f', stroked='f', **{'o:allowincell': 'f'})
    # Auto-fit shrinks the physical width to the longest line in Word, breaking
    # page-coordinate alignment. Fixed bounds plus zero insets preserve layout.
    textbox = _node('v:textbox', inset='0,0,0,0', style='mso-fit-shape-to-text:f')
    content = OxmlElement('w:txbxContent')
    # Each line remains a native paragraph; no HTML/XML supplied by the model.
    for line in block['text'].replace('\r\n', '\n').replace('\r', '\n').split('\n'):
        p = OxmlElement('w:p'); content.append(p)
        paragraph = Paragraph(p, anchor._parent)
        paragraph.add_run(line)
        _format_paragraph(paragraph, block)
    textbox.append(content); shape.append(textbox); pict.append(shape)
    anchor.add_run()._r.append(pict)


def _table(document, block, box):
    left, top, width, height = box
    rows = block['rows']; table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.autofit = False; table.alignment = WD_TABLE_ALIGNMENT.LEFT
    props = table._tbl.tblPr
    # Clear generated widths/layout before supplying the fixed physical grid.
    for tag in ('tblW', 'tblLayout', 'tblLook'):
        for old in props.findall(qn('w:' + tag)): props.remove(old)
    _w_property(props, 'tblW', w=round(width * 20), type='dxa')
    _w_property(props, 'tblLayout', type='fixed')
    _w_property(props, 'tblpPr', horzAnchor='page', vertAnchor='page', tblpX=round(left * 20),
                tblpY=round(top * 20), leftFromText='0', rightFromText='0', topFromText='0', bottomFromText='0')
    _w_property(props, 'tblOverlap', val='overlap')
    if block['rtl']: _w_property(props, 'bidiVisual', val='1')
    borders = _w_property(props, 'tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        _w_property(borders, edge, val='single', sz='4', color='D9D9D9')
    margins = _w_property(props, 'tblCellMar')
    for edge, twips in (('top', 35), ('bottom', 35), ('left', 45), ('right', 45)):
        _w_property(margins, edge, w=twips, type='dxa')
    _ordered(props, 'tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription tblPrChange')
    weights = block['column_widths']; total = sum(weights)
    widths = [width * weight / total for weight in weights]
    for column, points in zip(table.columns, widths): column.width = Pt(points)
    for ri, (row, cells) in enumerate(zip(table.rows, rows)):
        # At-least heights permit editing without silently clipping text.
        _w_property(row._tr.get_or_add_trPr(), 'trHeight', val=max(1, round(height * 20 / len(rows))), hRule='atLeast')
        _w_property(row._tr.get_or_add_trPr(), 'cantSplit')
        if block['header'] and ri == 0: _w_property(row._tr.get_or_add_trPr(), 'tblHeader')
        _ordered(row._tr.get_or_add_trPr(), 'cnfStyle divId gridBefore gridAfter wBefore wAfter cantSplit trHeight tblHeader tblCellSpacing jc hidden ins del trPrChange')
        for ci, (cell, text) in enumerate(zip(row.cells, cells)):
            cell.width = Pt(widths[ci]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = text
            cell_block = dict(block)
            if block['header'] and ri == 0:
                cell_block['bold'] = True
                _w_property(cell._tc.get_or_add_tcPr(), 'shd', fill='E9EDF2', val='clear')
            for paragraph in cell.paragraphs: _format_paragraph(paragraph, cell_block)
    for r0, c0, r1, c1 in block['merges']:
        table.cell(r0, c0).merge(table.cell(r1, c1))
    return table


def _load_page_image(source):
    if isinstance(source, Image.Image):
        if source.width * source.height > MAX_IMAGE_PIXELS:
            raise ValueError('page image exceeds pixel limit')
        return ImageOps.exif_transpose(source).convert('RGB')
    if not isinstance(source, (str, os.PathLike)):
        raise ValueError('page image must be a local path or PIL image')
    path = Path(source)
    if not path.is_file() or path.stat().st_size > MAX_IMAGE_BYTES:
        raise ValueError('missing or oversized page image')
    with Image.open(path) as image:
        if image.width * image.height > MAX_IMAGE_PIXELS:
            raise ValueError('page image exceeds pixel limit')
        return ImageOps.exif_transpose(image).convert('RGB')


def _picture(anchor, block, box, image, budget):
    left, top, width, height = block['bbox']
    pixels = (math.floor(left * image.width), math.floor(top * image.height),
              min(image.width, math.ceil((left + width) * image.width)),
              min(image.height, math.ceil((top + height) * image.height)))
    crop = image.crop(pixels)
    # The source need not be huge to support a small graphic; 220 dpi is ample.
    max_size = (max(1, math.ceil(box[2] / 72 * 220)), max(1, math.ceil(box[3] / 72 * 220)))
    crop.thumbnail(max_size, Image.Resampling.LANCZOS)
    budget['pixels'] += crop.width * crop.height
    if budget['pixels'] > MAX_CROP_PIXELS:
        crop.close(); raise ValueError('aggregate picture-crop pixel limit exceeded')
    stream = io.BytesIO(); crop.save(stream, format='PNG'); crop.close(); stream.seek(0)
    budget['bytes'] += stream.getbuffer().nbytes
    if budget['bytes'] > MAX_EMBEDDED_BYTES:
        stream.close(); raise ValueError('aggregate embedded picture byte limit exceeded')
    inline = anchor.add_run().add_picture(stream, width=Pt(box[2]), height=Pt(box[3]))._inline
    inline.docPr.set('descr', 'Picture crop from source page; surrounding text is editable')
    floating = _node('wp:anchor', distT='0', distB='0', distL='0', distR='0', simplePos='0',
                     relativeHeight='0', behindDoc='0', locked='0', layoutInCell='0', allowOverlap='1')
    floating.append(_node('wp:simplePos', x='0', y='0'))
    for direction, value in (('H', box[0]), ('V', box[1])):
        position = _node('wp:position' + direction, relativeFrom='page')
        offset = _node('wp:posOffset'); offset.text = str(round(value * 12700))
        position.append(offset); floating.append(position)
    extent = inline.find(qn('wp:extent')); inline.remove(extent); floating.append(extent)
    floating.append(_node('wp:wrapNone'))
    for child in list(inline): inline.remove(child); floating.append(child)
    inline.getparent().replace(inline, floating)


def _section(section, page):
    section.page_width = Pt(page['width_pt']); section.page_height = Pt(page['height_pt'])
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Pt(0)
    section.header_distance = section.footer_distance = Pt(0)
    section.gutter = Pt(0)
    for grid in section._sectPr.findall(qn('w:docGrid')): section._sectPr.remove(grid)


def build_docx(pages, output_path, page_images=None):
    """Atomically write a DOCX and return its Path; no OCR/network/COM execution.

    page_images is an index-aligned sequence or zero-based mapping of local
    image paths/PIL images; only pages with picture blocks need an image. The
    destination may be a .tmp staging file for an outer collision-safe writer.
    This function replaces an explicitly supplied existing destination; callers
    should use atomic_output for user-visible collision-safe publication.
    """
    pages = validate_pages(pages)
    output = Path(output_path).absolute()
    if output.suffix.lower() not in ('.docx', '.tmp') or not output.parent.is_dir() or output.is_dir():
        raise ValueError('output must be a .docx/.tmp file inside an existing directory')
    if page_images is None: page_images = {}
    if not isinstance(page_images, Mapping) and (not isinstance(page_images, Sequence) or isinstance(page_images, (str, bytes))):
        raise ValueError('page_images must be an index-aligned sequence or mapping')
    document = Document()
    normal = document.styles['Normal']; normal.font.name = 'Arial'; normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(0)
    document.core_properties.title = 'Editable reconstructed document'
    document.core_properties.author = 'Orbit Files'
    document.core_properties.last_modified_by = 'Orbit Files'
    total_pixels = shape_id = 0
    picture_budget = {'pixels': 0, 'bytes': 0}
    for pi, page in enumerate(pages):
        if pi: document.add_section(WD_SECTION_START.NEW_PAGE)
        _section(document.sections[-1], page)
        anchor = _anchor_paragraph(document)
        image = None
        try:
            if any(b['kind'] == 'picture' for b in page['blocks']):
                try: source = page_images[pi]
                except (KeyError, IndexError): raise ValueError(f'page_images[{pi}] is required for picture blocks') from None
                if not isinstance(source, Image.Image) and Path(source).absolute() == output:
                    raise ValueError('output cannot replace a source page image')
                image = _load_page_image(source)
                total_pixels += image.width * image.height
                if total_pixels > MAX_TOTAL_IMAGE_PIXELS:
                    raise ValueError('aggregate page-image pixel limit exceeded')
            for block in page['blocks']:
                left, top, width, height = block['bbox']
                box = (left * page['width_pt'], top * page['height_pt'],
                       width * page['width_pt'], height * page['height_pt'])
                shape_id += 1
                if block['kind'] == 'text': _textbox(anchor, block, box, shape_id)
                elif block['kind'] == 'table': _table(document, block, box)
                else: _picture(anchor, block, box, image, picture_budget)
        finally:
            if image is not None: image.close()
    # Keep section-break carrier paragraphs tiny; floating content does not need
    # body-flow placeholders and the last page has no trailing page break.
    for paragraph in document.paragraphs:
        paragraph.paragraph_format.space_after = paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.line_spacing = Pt(1)
        for run in paragraph.runs:
            if not run.text: run.font.size = Pt(1)
    fd, staging = tempfile.mkstemp(prefix='.orbit-docx-', suffix='.docx', dir=output.parent)
    os.close(fd)
    try:
        document.save(staging)
        if os.path.getsize(staging) > MAX_OUTPUT_BYTES:
            raise ValueError('generated DOCX exceeds output limit')
        os.replace(staging, output)
    finally:
        if os.path.exists(staging): os.unlink(staging)
    return output
